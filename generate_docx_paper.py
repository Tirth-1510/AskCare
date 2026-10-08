import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import re

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Page Margins: 1 inch (72 pt = 1440 dxa)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles setup
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(10.5)
    font_normal.color.rgb = RGBColor(45, 55, 72) # #2D3748

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run(
        "AskCare: A Resilient, Privacy-Preserving Clinical Query Resolution System "
        "Leveraging Hybrid Local-Edge Retrieval-Augmented Generation and Small Language Models"
    )
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 42, 77) # Deep Navy

    # Authors
    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_after = Pt(2)
    run_auth = p_auth.add_run("Patel Tirth, Pathan Kazmeenkhan, Patel Daksh\nGuided by: Dr. Vivek Shah")
    run_auth.font.size = Pt(11)
    run_auth.font.bold = True
    run_auth.font.color.rgb = RGBColor(43, 108, 176) # #2B6CB0

    # Affiliation
    p_aff = doc.add_paragraph()
    p_aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aff.paragraph_format.space_after = Pt(16)
    run_aff = p_aff.add_run(
        "Department of Computer Engineering and Information Technology\n"
        "Sankalchand Patel College of Engineering, Sankalchand Patel University, Visnagar, Gujarat, India"
    )
    run_aff.font.size = Pt(9.5)
    run_aff.font.italic = True
    run_aff.font.color.rgb = RGBColor(113, 128, 150) # Slate Grey

    # Horizontal Divider Table
    divider = doc.add_table(rows=1, cols=1)
    divider.alignment = WD_TABLE_ALIGNMENT.CENTER
    d_cell = divider.rows[0].cells[0]
    set_cell_background(d_cell, "1A365D")
    divider.rows[0].height = Pt(2)
    d_cell.paragraphs[0].text = ""

    # Abstract Callout Box
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    tbl_abs = doc.add_table(rows=1, cols=1)
    tbl_abs.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_abs.autofit = False
    tbl_abs.columns[0].width = Inches(6.5)
    cell_abs = tbl_abs.rows[0].cells[0]
    set_cell_background(cell_abs, "F7FAFC")
    set_cell_margins(cell_abs, top=140, bottom=140, left=200, right=200)

    p_abs = cell_abs.paragraphs[0]
    p_abs.paragraph_format.space_after = Pt(6)
    r_abs_title = p_abs.add_run("ABSTRACT\n")
    r_abs_title.font.bold = True
    r_abs_title.font.size = Pt(10.5)
    r_abs_title.font.color.rgb = RGBColor(26, 54, 93)

    r_abs_body = p_abs.add_run(
        "Modern healthcare informatics demands artificial intelligence systems that balance clinical factual "
        "accuracy, patient data privacy, and operational fault tolerance. While commercial cloud-hosted Large Language "
        "Models (LLMs) demonstrate high conversational fluency, their practical adoption in point-of-care clinical triage "
        "is severely impeded by data transmission risks under HIPAA/GDPR, prohibitive cloud operational expenses, high network "
        "latency, and susceptibility to catastrophic clinical hallucinations. This paper presents AskCare, an end-to-end, "
        "privacy-preserving, and edge-resilient clinical query resolution framework that synergizes task-specialized Small "
        "Language Models (SLMs) with a robust Hybrid Dense-Sparse Retrieval-Augmented Generation (RAG) architecture and dual-persistence "
        "storage.\n\n"
        "AskCare introduces three core engineering innovations: (1) a multi-stage Deterministic Clinical Domain Guard that "
        "screens input queries against medical lexical ontologies at zero inference cost, preventing non-medical prompt exploitation; "
        "(2) a Hybrid Dense-Sparse RAG Engine featuring an on-device WebAssembly transformer (all-MiniLM-L6-v2, 384-dimensional dense vectors) "
        "paired with an offline Normalized Polynomial Term-Frequency Hashing Vectorizer fallback, guaranteeing uninterrupted semantic retrieval "
        "even during total network or dependency partition; and (3) a Fault-Tolerant Dual-Persistence Database Tier that automatically fails over "
        "from MongoDB Atlas to an ACID-safe local JSON store without session loss. Extensive experimental verification confirms that AskCare "
        "operates with sub-100ms API routing, eliminates cloud-dependent data leakage through local Ollama execution of fine-tuned SLMs "
        "(including SmolLM2-Medical), and maintains 100% service uptime during complete remote database dropouts. This framework establishes an "
        "accessible, production-grade blueprint for secure, low-resource medical conversational AI in decentralized clinical facilities."
    )
    r_abs_body.font.size = Pt(9.5)
    r_abs_body.font.color.rgb = RGBColor(45, 55, 72)

    p_kw = cell_abs.add_paragraph()
    p_kw.paragraph_format.space_before = Pt(6)
    p_kw.paragraph_format.space_after = Pt(0)
    r_kw_title = p_kw.add_run("Keywords: ")
    r_kw_title.font.bold = True
    r_kw_title.font.size = Pt(9.5)
    r_kw_title.font.color.rgb = RGBColor(26, 54, 93)
    r_kw_body = p_kw.add_run(
        "Small Language Models (SLMs), Patient Query Resolution, Healthcare Informatics, Retrieval-Augmented Generation (RAG), "
        "Edge Computing, Data Privacy, Offline Vectorization, Ollama, Clinical Decision Support Systems."
    )
    r_kw_body.font.size = Pt(9.5)
    r_kw_body.font.color.rgb = RGBColor(74, 85, 104)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Read markdown and parse sections
    with open('AskCare_Research_Paper.md', 'r', encoding='utf-8') as f:
        md_text = f.read()

    # Split into lines and process
    lines = md_text.split('\n')
    idx = 0
    while idx < len(lines):
        line = lines[idx]

        # Skip headers / abstract already manually rendered
        if idx < 30 and ("# AskCare" in line or "### Abstract" in line or "Patel Tirth" in line or "**Keywords:**" in line):
            idx += 1
            continue

        # Major Headings
        if line.startswith('## '):
            heading_text = line.replace('## ', '').strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(heading_text)
            r.font.name = 'Calibri'
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(26, 54, 93) # #1A365D
            idx += 1
            continue

        # Sub-Headings
        if line.startswith('### '):
            sub_text = line.replace('### ', '').strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sub_text)
            r.font.name = 'Calibri'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(43, 108, 176) # #2B6CB0
            idx += 1
            continue

        # Sub-Sub-Headings
        if line.startswith('#### '):
            sub_sub = line.replace('#### ', '').strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sub_sub)
            r.font.name = 'Calibri'
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(74, 85, 104)
            idx += 1
            continue

        # Markdown Tables
        if line.startswith('| ') and idx + 1 < len(lines) and lines[idx+1].startswith('| :---'):
            table_lines = []
            while idx < len(lines) and lines[idx].startswith('|'):
                table_lines.append(lines[idx])
                idx += 1

            headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
            rows = []
            for r_line in table_lines[2:]:
                rows.append([c.strip() for c in r_line.split('|')[1:-1]])

            table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            # Style header row
            hdr_cells = table.rows[0].cells
            for col_idx, h_text in enumerate(headers):
                cell = hdr_cells[col_idx]
                cell.text = h_text
                set_cell_background(cell, "1A365D")
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.bold = True
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(255, 255, 255)

            # Style data rows
            for row_idx, r_data in enumerate(rows):
                row_cells = table.rows[row_idx + 1].cells
                bg_color = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
                for col_idx, cell_data in enumerate(r_data):
                    if col_idx < len(row_cells):
                        cell = row_cells[col_idx]
                        clean_data = cell_data.replace('**', '').replace('*', '')
                        cell.text = clean_data
                        set_cell_background(cell, bg_color)
                        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                        for p in cell.paragraphs:
                            if col_idx > 0:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for r in p.runs:
                                r.font.size = Pt(8.5)
                                if "PASS" in clean_data:
                                    r.font.bold = True
                                    r.font.color.rgb = RGBColor(40, 167, 69)

            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # Code Blocks / Architecture Diagrams
        if line.startswith('```'):
            code_lines = []
            idx += 1
            while idx < len(lines) and not lines[idx].startswith('```'):
                code_lines.append(lines[idx])
                idx += 1
            idx += 1 # skip ending ```

            code_text = '\n'.join(code_lines)
            tbl_code = doc.add_table(rows=1, cols=1)
            tbl_code.alignment = WD_TABLE_ALIGNMENT.CENTER
            c_cell = tbl_code.rows[0].cells[0]
            set_cell_background(c_cell, "0F172A") # Dark background
            set_cell_margins(c_cell, top=120, bottom=120, left=150, right=150)
            p_code = c_cell.paragraphs[0]
            p_code.paragraph_format.space_after = Pt(0)
            r_code = p_code.add_run(code_text)
            r_code.font.name = 'Consolas'
            r_code.font.size = Pt(8.0)
            r_code.font.color.rgb = RGBColor(212, 255, 0) # Neon/Monospace text
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # Lists
        if line.startswith('* ') or line.startswith('- ') or re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Bullet' if (line.startswith('* ') or line.startswith('- ')) else 'List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            content = re.sub(r'^(\*|-|\d+\.)\s+', '', line)
            
            # Simple markdown bold parser
            parts = re.split(r'(\*\*.*?\*\*)', content)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
            idx += 1
            continue

        # Regular Paragraph
        clean_line = line.strip()
        if clean_line and clean_line != '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', clean_line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    r = p.add_run(part[1:-1])
                    r.font.italic = True
                else:
                    p.add_run(part)

        idx += 1

    doc.save('AskCare_Production_Research_Paper.docx')
    print("Document successfully created: AskCare_Production_Research_Paper.docx")

if __name__ == '__main__':
    create_document()
